import html
import re
from dataclasses import dataclass
from pathlib import Path
from typing import List

from openpyxl import load_workbook
from openpyxl.utils import get_column_letter

# Control characters that break TSV/clipboard paste into Excel (tab splits columns).
_CONTROL_CHAR_RE = re.compile(r"[\x00-\x08\x0b\x0c\x0e-\x1f]")


@dataclass
class ExportData:
    """Data rows from exported Finish Status Report (header excluded)."""

    rows: List[List[str]]
    row_count: int
    col_count: int
    tsv: str

    @property
    def clear_through_row(self) -> int:
        """Last row to clear in SharePoint sheet (header is row 1)."""
        return max(self.row_count + 500, 5000)


def sanitize_cell_value(cell) -> str:
    """Normalize a cell for SharePoint paste. Tabs/newlines split columns in Excel TSV paste."""
    if cell is None:
        return ""
    text = str(cell)
    text = text.replace("\r\n", " ").replace("\n", " ").replace("\r", " ")
    text = text.replace("\t", " ")
    text = _CONTROL_CHAR_RE.sub("", text)
    return " ".join(text.split())


def read_export_data(source_path: Path) -> ExportData:
    """Read all data rows below the header, skipping rows without valid IDs."""
    wb = load_workbook(source_path, data_only=True)
    ws = wb.active
    max_row = ws.max_row or 1
    max_col = ws.max_column or 1

    if max_row < 2:
        wb.close()
        return ExportData(rows=[], row_count=0, col_count=max_col, tsv="")

    data_rows: List[List[str]] = []
    for row in ws.iter_rows(min_row=2, max_row=max_row, max_col=max_col, values_only=True):
        # Get first cell value
        first_cell = row[0] if row else None
        first_cell_str = str(first_cell).strip() if first_cell else ""

        # Only include rows where first column is a number (ID)
        if first_cell_str and first_cell_str.isdigit():
            data_rows.append([sanitize_cell_value(cell) for cell in row])

    wb.close()

    col_count = max((len(r) for r in data_rows), default=max_col)
    tsv = values_to_tsv(data_rows)
    return ExportData(
        rows=data_rows,
        row_count=len(data_rows),
        col_count=col_count,
        tsv=tsv,
    )


def read_values_without_header(source_path: Path) -> List[List[str]]:
    return read_export_data(source_path).rows


def values_to_tsv(values: List[List[str]]) -> str:
    if not values:
        return ""
    col_count = max(len(row) for row in values)
    lines: list[str] = []
    for row in values:
        cells = [sanitize_cell_value(cell) for cell in row]
        while len(cells) < col_count:
            cells.append("")
        lines.append("\t".join(cells))
    return "\n".join(lines)


def rows_to_html_table(values: List[List[str]]) -> str:
    """HTML table paste keeps each value in one cell even if it contained tabs."""
    if not values:
        return "<table></table>"
    col_count = max(len(row) for row in values)
    parts = ["<table>"]
    for row in values:
        parts.append("<tr>")
        cells = [sanitize_cell_value(cell) for cell in row]
        while len(cells) < col_count:
            cells.append("")
        for cell in cells:
            parts.append(f"<td>{html.escape(cell)}</td>")
        parts.append("</tr>")
    parts.append("</table>")
    return "\n".join(parts)


def clear_range_address(col_count: int, clear_through_row: int) -> str:
    col = get_column_letter(max(col_count, 1))
    return f"A2:{col}{clear_through_row}"
